#!/usr/bin/env python3
"""Build a read-only HarmonyOS legal index from the complete source corpus."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree

import pdfplumber
from openpyxl import load_workbook
from pypdf import PdfReader


CORE_DIRECTORIES = ("法律", "行政法规", "规章及法律规范")
ARTICLE_PATTERN = re.compile(
    r"^第[零〇一二三四五六七八九十百千万\d]+条"
    r"(?:之[零〇一二三四五六七八九十百千万\d]+)?",
    re.MULTILINE,
)
PAGE_NUMBER_PATTERN = re.compile(r"^\s*\d+\s*/\s*\d+\s*$")
EFFECTIVE_DATE_PATTERN = re.compile(r"_(\d{4}\.\d{2}\.\d{2})生效")
STATUS_PATTERN = re.compile(r"时\s*效\s*性\s*[：:]\s*([^\n]+)")
DOWNLOAD_SUFFIX_PATTERN = re.compile(
    r"(?:_\d{4}\.\d{2}\.\d{2}生效)?_\d{8}下载(?:\s*\(\d+\))?$"
)
KNOWN_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".wps", ".xls", ".xlsx", ".ofd", ".zip", ".rar",
    ".jpg", ".jpeg", ".png",
}
MAX_ARCHIVE_BYTES = 250 * 1024 * 1024


@dataclass(frozen=True)
class Extraction:
    text: str
    page_count: int
    file_format: str


class VisibleHtmlParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.ignored_depth = 0

    def handle_starttag(
        self, tag: str, attrs: list[tuple[str, str | None]]
    ) -> None:
        if tag in {"script", "style"}:
            self.ignored_depth += 1
        elif self.ignored_depth == 0 and tag in {
            "br", "p", "div", "tr", "td", "th", "li", "h1", "h2", "h3",
        }:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self.ignored_depth > 0:
            self.ignored_depth -= 1
        elif self.ignored_depth == 0 and tag in {
            "p", "div", "tr", "td", "th", "li", "h1", "h2", "h3",
        }:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self.ignored_depth == 0:
            self.parts.append(data)


@dataclass(frozen=True)
class LawDocument:
    documentId: str
    title: str
    category: str
    subcategory: str
    region: str
    source: str
    sourcePath: str
    fileFormat: str
    extractionStatus: str
    extractionNote: str
    status: str
    effectiveDate: str
    pageCount: int
    contentHash: str


@dataclass(frozen=True)
class LawChunk:
    id: str
    documentId: str
    title: str
    article: str
    category: str
    subcategory: str
    region: str
    source: str
    status: str
    effectiveDate: str
    text: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert the complete labor-law corpus to a compact JSON index."
    )
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    parser.add_argument(
        "--scope",
        choices=("all", "core"),
        default="all",
        help="all: every source file; core: the original 84 national PDF files.",
    )
    parser.add_argument("--max-chars", type=int, default=1200)
    parser.add_argument("--overlap-chars", type=int, default=120)
    parser.add_argument("--soffice", type=Path)
    parser.add_argument(
        "--strict",
        action="store_true",
        help="Return a non-zero status if any source file has no extractable text.",
    )
    return parser.parse_args()


def source_files(source_root: Path, scope: str) -> list[Path]:
    if not source_root.is_dir():
        raise FileNotFoundError(f"Missing source root: {source_root}")
    if scope == "core":
        files: list[Path] = []
        for directory in CORE_DIRECTORIES:
            target = source_root / directory
            if not target.is_dir():
                raise FileNotFoundError(f"Missing target directory: {target}")
            files.extend(target.rglob("*.pdf"))
    else:
        files = [
            path
            for path in source_root.rglob("*")
            if path.is_file()
            and path.name != ".DS_Store"
            and not path.name.startswith("._")
        ]
    return sorted(files, key=lambda item: item.relative_to(source_root).as_posix())


def clean_title(path: Path) -> str:
    name = path.name
    if path.suffix.lower() in KNOWN_EXTENSIONS:
        name = name[: -len(path.suffix)]
    return DOWNLOAD_SUFFIX_PATTERN.sub("", name).strip()


def normalize_text(text: str, title: str = "") -> str:
    kept_lines: list[str] = []
    compact_title = re.sub(r"\s+", "", title)
    for raw_line in text.replace("\u00a0", " ").replace("\r", "\n").splitlines():
        line = raw_line.strip()
        if not line or PAGE_NUMBER_PATTERN.fullmatch(line):
            continue
        if "威科先行" in line and "法律信息库" in line:
            continue
        if "扫一扫，手机阅读更方便" in line:
            continue
        line = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", line)
        line = re.sub(r"[ \t]+", " ", line)
        if compact_title and re.sub(r"\s+", "", line) == compact_title:
            continue
        kept_lines.append(line)
    return "\n".join(kept_lines)


def extract_pdf(path: Path) -> Extraction:
    title = clean_title(path)
    page_texts: list[str] = []
    page_count = 0
    first_error: Exception | None = None
    try:
        with pdfplumber.open(path) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                page_texts.append(normalize_text(page.extract_text() or "", title))
    except Exception as exc:
        first_error = exc

    text = "\n".join(part for part in page_texts if part)
    if len(text.strip()) >= 40:
        return Extraction(text, page_count, "PDF")

    try:
        reader = PdfReader(path)
        page_count = len(reader.pages)
        text = "\n".join(
            normalize_text(page.extract_text() or "", title) for page in reader.pages
        )
    except Exception:
        if first_error is not None:
            raise first_error
        raise
    return Extraction(text, page_count, "PDF")


def _xml_text(stream: object) -> str:
    parts: list[str] = []
    for _event, element in ElementTree.iterparse(stream, events=("end",)):
        local_name = element.tag.rsplit("}", 1)[-1]
        if local_name in {"t", "TextCode"} and element.text:
            parts.append(element.text)
        elif local_name == "tab":
            parts.append("\t")
        elif local_name in {"br", "p", "tr"}:
            parts.append("\n")
        element.clear()
    return "".join(parts)


def extract_docx(path: Path) -> Extraction:
    text_parts: list[str] = []
    with zipfile.ZipFile(path) as archive:
        candidates = [
            name
            for name in archive.namelist()
            if name == "word/document.xml"
            or re.fullmatch(
                r"word/(?:header\d+|footer\d+|footnotes|endnotes|comments)\.xml",
                name,
            )
        ]
        if "word/document.xml" not in candidates:
            raise ValueError("not a Word OOXML document")
        for name in candidates:
            with archive.open(name) as stream:
                text_parts.append(_xml_text(stream))
    return Extraction(normalize_text("\n".join(text_parts), clean_title(path)), 1, "DOCX")


def extract_xlsx(path: Path) -> Extraction:
    lines: list[str] = []
    with path.open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=True)
        sheet_count = len(workbook.worksheets)
        for sheet in workbook.worksheets:
            lines.append(f"工作表：{sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    lines.append("\t".join(values))
        workbook.close()
    return Extraction(normalize_text("\n".join(lines), clean_title(path)), sheet_count, "XLSX")


def extract_ofd(path: Path) -> Extraction:
    text_parts: list[str] = []
    page_count = 0
    with zipfile.ZipFile(path) as archive:
        page_names = sorted(
            name
            for name in archive.namelist()
            if name.endswith("/Content.xml") and "/Pages/" in name
        )
        if not page_names:
            raise ValueError("OFD has no page content")
        page_count = len(page_names)
        for name in page_names:
            with archive.open(name) as stream:
                text_parts.append(_xml_text(stream))
    return Extraction(normalize_text("\n".join(text_parts), clean_title(path)), page_count, "OFD")


def extract_html(path: Path) -> Extraction:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    if "系统出现异常" in raw or "无权限下载" in raw:
        raise ValueError("downloaded file is an HTML error page")
    parser = VisibleHtmlParser()
    parser.feed(raw)
    text = normalize_text("".join(parser.parts), clean_title(path))
    file_format = "HTML/XLS" if "urn:schemas-microsoft-com:office:excel" in raw else "HTML"
    return Extraction(text, 1, file_format)


def run_soffice(source: Path, target_format: str, soffice: Path) -> Path:
    with tempfile.TemporaryDirectory(prefix="law-kb-office-profile-") as profile_dir:
        command = [
            str(soffice),
            "--headless",
            f"-env:UserInstallation=file://{profile_dir}",
            "--convert-to",
            target_format,
            "--outdir",
            str(source.parent),
            str(source),
        ]
        result = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
    target = source.with_suffix(f".{target_format}")
    if result.returncode != 0 or not target.exists():
        message = (result.stderr or result.stdout or "conversion failed").strip()
        raise RuntimeError(message)
    return target


def extract_legacy_office(path: Path, soffice: Path | None) -> Extraction:
    if soffice is None or not soffice.exists():
        raise RuntimeError("LibreOffice/soffice is unavailable")
    errors: list[str] = []
    with tempfile.TemporaryDirectory(prefix="law-kb-office-") as temp_dir:
        temp_root = Path(temp_dir)
        for source_suffix, target_format, extractor in (
            (".doc", "docx", extract_docx),
            (".xls", "xlsx", extract_xlsx),
        ):
            candidate = temp_root / f"source{source_suffix}"
            shutil.copyfile(path, candidate)
            try:
                converted = run_soffice(candidate, target_format, soffice)
                result = extractor(converted)
                if len(result.text.strip()) >= 40:
                    return Extraction(
                        result.text,
                        result.page_count,
                        "WPS/DOC" if target_format == "docx" else "XLS",
                    )
            except Exception as exc:
                errors.append(f"{target_format}: {exc}")
    raise RuntimeError("; ".join(errors) or "legacy Office conversion produced no text")


def detect_kind(path: Path) -> str:
    with path.open("rb") as source:
        header = source.read(256)
    stripped_header = header.lstrip().lower()
    if stripped_header.startswith(b"<html") or stripped_header.startswith(b"<!doctype"):
        return "html"
    if header.startswith(b"%PDF"):
        return "pdf"
    if header.startswith(b"\xd0\xcf\x11\xe0"):
        return "legacy"
    if header.startswith(b"PK"):
        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
            if "word/document.xml" in names:
                return "docx"
            if "xl/workbook.xml" in names:
                return "xlsx"
            if "OFD.xml" in names:
                return "ofd"
            return "archive"
        except zipfile.BadZipFile:
            return "broken_archive"
    suffix = path.suffix.lower()
    if suffix in {".jpg", ".jpeg", ".png"}:
        return "image"
    if suffix == ".rar":
        return "rar"
    return "unknown"


def extract_archive(path: Path, soffice: Path | None) -> Extraction:
    sections: list[str] = []
    total_pages = 0
    extracted_bytes = 0
    with zipfile.ZipFile(path) as archive, tempfile.TemporaryDirectory(
        prefix="law-kb-archive-"
    ) as temp_dir:
        for index, info in enumerate(archive.infolist(), start=1):
            if info.is_dir():
                continue
            extracted_bytes += info.file_size
            if extracted_bytes > MAX_ARCHIVE_BYTES:
                raise ValueError("archive expands beyond the 250 MB safety limit")
            safe_name = Path(info.filename).name or f"entry-{index}"
            temp_path = Path(temp_dir) / f"{index:04d}-{safe_name}"
            with archive.open(info) as source, temp_path.open("wb") as target:
                shutil.copyfileobj(source, target)
            try:
                result = extract_path(temp_path, soffice, allow_archive=False)
            except Exception:
                continue
            if len(result.text.strip()) < 40:
                continue
            sections.append(f"压缩包内文件：{safe_name}\n{result.text}")
            total_pages += max(result.page_count, 1)
    if not sections:
        raise ValueError("archive contains no supported text document")
    return Extraction(normalize_text("\n\n".join(sections), clean_title(path)), total_pages, "ZIP")


def extract_path(path: Path, soffice: Path | None, allow_archive: bool = True) -> Extraction:
    kind = detect_kind(path)
    if kind == "pdf":
        return extract_pdf(path)
    if kind == "docx":
        return extract_docx(path)
    if kind == "xlsx":
        return extract_xlsx(path)
    if kind == "ofd":
        return extract_ofd(path)
    if kind == "html":
        return extract_html(path)
    if kind == "legacy":
        return extract_legacy_office(path, soffice)
    if kind == "archive" and allow_archive:
        return extract_archive(path, soffice)
    if kind == "image":
        raise ValueError("image OCR is intentionally disabled")
    if kind == "rar":
        raise ValueError("RAR extraction is unavailable")
    if kind == "broken_archive":
        raise ValueError("archive is damaged or incomplete")
    raise ValueError(f"unsupported file type: {path.suffix or 'no extension'}")


def effective_date(path: Path, text: str) -> str:
    filename_match = EFFECTIVE_DATE_PATTERN.search(path.name)
    if filename_match:
        return filename_match.group(1)
    text_match = re.search(
        r"生\s*效\s*日\s*期\s*[：:]\s*(\d{4}[.\-/]\d{2}[.\-/]\d{2})",
        text[:1500],
    )
    return text_match.group(1) if text_match else ""


def validity_status(path: Path, text: str) -> str:
    match = STATUS_PATTERN.search(text[:1800])
    if match:
        return match.group(1).strip()
    if "失效" in path.name or "废止" in path.name:
        return "文件名标注失效或废止"
    return "未标注"


def split_long_text(text: str, max_chars: int, overlap_chars: int) -> list[str]:
    if len(text) <= max_chars:
        return [text]
    parts: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            candidates = [
                text.rfind("\n", start + max_chars // 2, end),
                text.rfind("。", start + max_chars // 2, end),
                text.rfind("；", start + max_chars // 2, end),
            ]
            boundary = max(candidates)
            if boundary > start:
                end = boundary + 1
        part = text[start:end].strip()
        if part:
            parts.append(part)
        if end >= len(text):
            break
        start = max(end - overlap_chars, start + 1)
    return parts


def article_sections(text: str) -> list[tuple[str, str]]:
    matches = []
    cross_reference_prefixes = ("、", "，", ",", "。", "；", ";", "）", ")", "和", "及", "至", "到")
    for match in ARTICLE_PATTERN.finditer(text):
        suffix = text[match.end() : match.end() + 1]
        if suffix.startswith(cross_reference_prefixes):
            continue
        matches.append(match)
    if not matches:
        return [("全文", text)]

    sections: list[tuple[str, str]] = []
    preamble = text[: matches[0].start()].strip()
    if len(preamble) >= 80:
        sections.append(("前言", preamble))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        sections.append((match.group(0), text[match.start() : end].strip()))
    return sections


def build_chunks(
    document: LawDocument,
    text: str,
    max_chars: int,
    overlap_chars: int,
) -> list[LawChunk]:
    chunks: list[LawChunk] = []
    chunk_number = 1
    for article, section in article_sections(text):
        for part_index, part in enumerate(
            split_long_text(section, max_chars, overlap_chars), start=1
        ):
            display_article = article
            if len(section) > max_chars:
                display_article = f"{article}（片段{part_index}）"
            chunks.append(
                LawChunk(
                    id=f"{document.documentId}-{chunk_number:04d}",
                    documentId=document.documentId,
                    title=document.title,
                    article=display_article,
                    category=document.category,
                    subcategory=document.subcategory,
                    region=document.region,
                    source=document.source,
                    status=document.status,
                    effectiveDate=document.effectiveDate,
                    text=part,
                )
            )
            chunk_number += 1
    return chunks


def document_location(relative: Path) -> tuple[str, str, str]:
    category = relative.parts[0] if relative.parts else "未分类"
    subcategory = "/".join(relative.parts[1:-1])
    region = "全国"
    if category.startswith("地方法律") and len(relative.parts) > 1:
        region = relative.parts[1]
    return category, subcategory, region


def iter_documents(
    source_root: Path,
    files: Iterable[Path],
    max_chars: int,
    overlap_chars: int,
    soffice: Path | None,
) -> tuple[list[LawDocument], list[LawChunk], list[str]]:
    documents: list[LawDocument] = []
    chunks: list[LawChunk] = []
    failures: list[str] = []
    file_list = list(files)
    for document_number, path in enumerate(file_list, start=1):
        relative = path.relative_to(source_root)
        category, subcategory, region = document_location(relative)
        document_id = f"LAW-{document_number:04d}"
        title = clean_title(path)
        try:
            result = extract_path(path, soffice)
            if len(result.text.strip()) < 40:
                raise ValueError("extracted text is empty or too short")
            digest = hashlib.sha256(result.text.encode("utf-8")).hexdigest()
            document = LawDocument(
                documentId=document_id,
                title=title,
                category=category,
                subcategory=subcategory,
                region=region,
                source=path.name,
                sourcePath=relative.as_posix(),
                fileFormat=result.file_format,
                extractionStatus="正文已提取",
                extractionNote="",
                status=validity_status(path, result.text),
                effectiveDate=effective_date(path, result.text),
                pageCount=result.page_count,
                contentHash=digest,
            )
            documents.append(document)
            chunks.extend(build_chunks(document, result.text, max_chars, overlap_chars))
        except Exception as exc:
            note = f"{type(exc).__name__}: {exc}"
            failures.append(f"{relative.as_posix()}: {note}")
            documents.append(
                LawDocument(
                    documentId=document_id,
                    title=title,
                    category=category,
                    subcategory=subcategory,
                    region=region,
                    source=path.name,
                    sourcePath=relative.as_posix(),
                    fileFormat=detect_kind(path).upper(),
                    extractionStatus="仅收录文件信息",
                    extractionNote=note,
                    status=validity_status(path, ""),
                    effectiveDate=effective_date(path, ""),
                    pageCount=0,
                    contentHash="",
                )
            )
        if document_number % 50 == 0 or document_number == len(file_list):
            print(
                f"[PROGRESS] {document_number}/{len(file_list)} files, "
                f"{len(chunks)} chunks, {len(failures)} metadata-only",
                flush=True,
            )
    return documents, chunks, failures


def main() -> int:
    args = parse_args()
    files = source_files(args.source_root, args.scope)
    soffice = args.soffice or (
        Path(shutil.which("soffice")) if shutil.which("soffice") else None
    )
    print(
        f"[INFO] scope={args.scope}, source files={len(files)}, "
        f"soffice={soffice or 'unavailable'}",
        flush=True,
    )
    documents, chunks, failures = iter_documents(
        args.source_root,
        files,
        args.max_chars,
        args.overlap_chars,
        soffice,
    )
    searchable_document_count = sum(
        item.extractionStatus == "正文已提取" for item in documents
    )
    payload = {
        "version": 3,
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "sourceFileCount": len(files),
        "documentCount": len(documents),
        "searchableDocumentCount": searchable_document_count,
        "metadataOnlyCount": len(failures),
        "chunkCount": len(chunks),
        "documents": [
            {
                "documentId": item.documentId,
                "title": item.title,
                "category": item.category,
                "subcategory": item.subcategory,
                "region": item.region,
                "source": item.source,
                "status": item.status,
                "effectiveDate": item.effectiveDate,
            }
            for item in documents
        ],
        "chunks": [
            {
                "id": item.id,
                "documentId": item.documentId,
                "article": item.article,
                "text": item.text,
            }
            for item in chunks
        ],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )

    category_counts = Counter(document.category for document in documents)
    format_counts = Counter(document.fileFormat for document in documents)
    report_lines = [
        "# 本地法律索引构建报告",
        "",
        f"- 源文件总数：{len(files)}",
        f"- 已收录文档记录：{len(documents)}",
        f"- 正文可检索文档：{searchable_document_count}",
        f"- 仅收录文件信息：{len(failures)}",
        f"- 条文/文本块数：{len(chunks)}",
        f"- 索引大小：{args.output.stat().st_size} bytes",
        "",
        "## 分类",
        "",
    ]
    report_lines.extend(
        f"- {category}：{count}" for category, count in sorted(category_counts.items())
    )
    report_lines.extend(["", "## 提取格式", ""])
    report_lines.extend(
        f"- {file_format}：{count}"
        for file_format, count in sorted(format_counts.items())
    )
    report_lines.extend(
        [
            "",
            "## 说明",
            "",
            "本工具只做文字提取与机械切分，不进行去重、失效清理、版本合并或法律效力判断。",
            "图片未执行 OCR；无法读取的压缩包或特殊格式只保留文件信息，不作为法规正文参与问答检索。",
        ]
    )
    if failures:
        report_lines.extend(["", "## 仅收录文件信息的文件", ""])
        report_lines.extend(f"- {failure}" for failure in failures)
    report = "\n".join(report_lines) + "\n"
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(report, encoding="utf-8")
    print(report, end="")
    if args.strict and failures:
        return 2
    if args.scope == "core" and (len(files) != 84 or searchable_document_count != 84):
        print(
            f"[ERROR] Core corpus mismatch: files={len(files)}, "
            f"searchable={searchable_document_count}, expected=84"
        )
        return 3
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
